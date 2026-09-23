"""Siapkan lembar anotasi manusia: 300 ulasan test untuk dibaca ulang oleh tim.

Kenapa perlu: label di dataset berasal dari bintang, dan pembuat dataset sendiri menulis bahwa
tidak ada anotator kedua untuk satu ulasan. Dengan membaca ulang 300 ulasan tanpa melihat
bintangnya, tim bisa mengukur dua hal: seberapa sering pembaca manusia sepakat satu sama lain,
dan seberapa sering label bintang berbeda dari pembacaan manusia.

Cara pakai:
1. Jalankan skrip ini. Lima berkas Excel muncul di ml/anotasi/lembar/.
2. Tiap anggota tim mengisi berkasnya sendiri, tanpa berdiskusi dan tanpa membuka folder kunci.
3. Berkas yang sudah terisi dikumpulkan ke ml/anotasi/hasil/ dengan nama tetap.
4. Jalankan ml/scripts/15_anotasi_hitung.py.

Output: ml/anotasi/lembar/anotator_N.xlsx, ml/anotasi/kunci/kunci.csv
"""

from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "ml"))
from suara_ml import prediksi  # noqa: E402
from suara_ml.textnorm import mask_personal  # noqa: E402

sys.path.insert(0, str(ROOT / "ml" / "scripts"))
from importlib import import_module  # noqa: E402

extras = import_module("09_web_extras")

LABELS = ["negative", "neutral", "positive"]
PILIHAN = ["negatif", "netral", "positif", "tidak jelas"]
PER_KELAS = 100
ANOTATOR = 5
SEED = 7
NAMA_APP = {
    "mobileJKN": "Mobile JKN",
    "JMO": "JMO",
    "satusehat": "SatuSehat",
    "pertamina": "MyPertamina",
    "KAI": "KAI Access",
    "BMKG": "Info BMKG",
}
PANDUAN = [
    ("negatif", "Penulis kecewa, mengeluh, melaporkan kerusakan, atau menyindir layanan."),
    ("netral", "Pertanyaan, saran, informasi, atau pujian dan keluhan yang seimbang."),
    ("positif", "Penulis puas, memuji, atau berterima kasih."),
    ("tidak jelas", "Bukan bahasa yang bisa dipahami, kosong, atau tidak berhubungan dengan aplikasi."),
]


def contoh() -> pd.DataFrame:
    prep = pd.read_parquet(ROOT / "data" / "igar_prepared.parquet")
    te = prep[prep["split"] == "test"].copy()
    t = te["text"].astype(str)
    layak = t.str.len().between(15, 220) & ~t.str.contains(r"@|http|www\.", regex=True) & ~t.str.contains(extras.SAPAAN_NAMA) & ~t.str.contains(extras.RUSAK)
    te = te[layak]
    bagian = [te[te["label"] == lab].sample(n=PER_KELAS, random_state=SEED) for lab in LABELS]
    s = pd.concat(bagian).sample(frac=1, random_state=SEED).reset_index(drop=True)
    s["ulasan"] = [mask_personal(" ".join(str(x).split())) for x in s["text"]]
    s["no"] = range(1, len(s) + 1)
    return s


def tulis_lembar(s: pd.DataFrame, nomor: int, tujuan: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "anotasi"
    ws.append(["no", "aplikasi", "ulasan", "label", "catatan"])
    for sel in ws[1]:
        sel.font = Font(bold=True, color="FFFFFF")
        sel.fill = PatternFill("solid", fgColor="B0151F")
    for _, r in s.iterrows():
        ws.append([int(r["no"]), NAMA_APP.get(r["app"], r["app"]), r["ulasan"], "", ""])
    dv = DataValidation(type="list", formula1='"' + ",".join(PILIHAN) + '"', allow_blank=True, showDropDown=False)
    dv.error = "Pilih salah satu dari daftar."
    ws.add_data_validation(dv)
    dv.add(f"D2:D{len(s) + 1}")
    lebar = {"A": 6, "B": 14, "C": 90, "D": 14, "E": 28}
    for kolom, w in lebar.items():
        ws.column_dimensions[kolom].width = w
    for baris in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        baris[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    p = wb.create_sheet("panduan")
    p.append([f"Lembar anotasi {nomor} dari {ANOTATOR}"])
    p["A1"].font = Font(bold=True, size=14)
    p.append([])
    for baris in [
        "Baca tiap ulasan lalu pilih satu label di kolom label.",
        "Nilai isi tulisannya saja. Jangan menebak berapa bintang yang diberi penulisnya.",
        "Isi sendiri tanpa berdiskusi dengan anggota lain, supaya tingkat kesepakatan berarti.",
        "Kolom catatan boleh dikosongkan. Isi kalau ulasannya membingungkan.",
        "Jangan membuka folder ml/anotasi/kunci sebelum semua baris terisi.",
        "",
        "Arti tiap label:",
    ]:
        p.append([baris])
    for label, arti in PANDUAN:
        p.append([label, arti])
    p.column_dimensions["A"].width = 16
    p.column_dimensions["B"].width = 80
    for baris in p.iter_rows(min_col=2, max_col=2):
        baris[0].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(tujuan)


def main() -> None:
    s = contoh()
    lembar = ROOT / "ml" / "anotasi" / "lembar"
    kunci = ROOT / "ml" / "anotasi" / "kunci"
    lembar.mkdir(parents=True, exist_ok=True)
    kunci.mkdir(parents=True, exist_ok=True)

    for i in range(1, ANOTATOR + 1):
        tulis_lembar(s, i, lembar / f"anotator_{i}.xlsx")

    jawab = prediksi.web_models(s["ulasan"].tolist())
    if prediksi.indobert_available():
        jawab["indobert"] = prediksi.indobert(s["ulasan"].tolist())
    kolom = {"no": s["no"], "app": s["app"], "ulasan": s["ulasan"], "label_bintang": s["label"], "bintang_rata": s["score_mean"].round(2), "n_ulasan_sama": s["n"]}
    kolom.update(jawab)
    pd.DataFrame(kolom).to_csv(kunci / "kunci.csv", index=False, encoding="utf-8", lineterminator="\n")

    print(f"{len(s)} ulasan, {ANOTATOR} lembar di {lembar}")
    print("sebaran label bintang:", s["label"].value_counts().to_dict())
    print("model di kunci:", list(jawab))


if __name__ == "__main__":
    main()
